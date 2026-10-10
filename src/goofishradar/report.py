"""每日报告生成（T07 / FR-7 P0）：HTML + Excel(4 Sheet)，归档 reports/YYYY-MM-DD/。

- 口径说明固定出现在 HTML 顶部（T07 验收标准 3）；
- 三榜数据来自 T06 评分结果（同一数据源，保证与库内评分一致）；
- 评分结果标准归档到 data/parsed/{date}/_rankings.json，供次日"与昨日对比"；
- 报告目录按 retain_days 自动清理过期。
"""
from __future__ import annotations

import json
import re
import shutil
from datetime import date, timedelta
from pathlib import Path

from jinja2 import Environment, PackageLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import Settings
from .ranking import rankings_to_dict, score_day_from_store
from .storage import Store

_ENV = Environment(loader=PackageLoader("goofishradar", "templates"),
                   autoescape=select_autoescape(["html", "j2"]))

_XLSX_HEADERS = ["排名", "商品ID", "标题", "价格", "想要数", "日增(1d)", "浏览量",
                 "评分", "卖家", "所在地", "状态", "来源关键词", "链接"]
_COL_WIDTHS = [6, 16, 46, 9, 9, 9, 9, 9, 14, 10, 8, 14, 40]


# ---- 与昨日对比 ----

def compare_with_yesterday(store: Store, date_str: str, today_ids: list[str],
                           data_dir: Path) -> dict:
    """与最近一次既往榜单存档对比（向前最多回溯 7 天，兼容中断日）。"""
    parsed_root = data_dir / "parsed"
    archive_file, compared_date = None, None
    for i in range(1, 8):
        d = (date.fromisoformat(date_str) - timedelta(days=i)).isoformat()
        f = parsed_root / d / "_rankings.json"
        if f.exists():
            archive_file, compared_date = f, d
            break
    if archive_file is None:
        return {"available": False,
                "note": "无既往榜单存档，暂无可比数据（对比自第二个报告日起可用）"}
    yday_ids = [it["item_id"]
                for it in json.loads(archive_file.read_text(encoding="utf-8"))["overall"]]
    today_set, yday_set = set(today_ids), set(yday_ids)
    new_ids = [i for i in today_ids if i not in yday_set][:20]
    dropped_ids = [i for i in yday_ids if i not in today_set][:20]

    def info(ids: list[str]) -> list[dict]:
        out = []
        for iid in ids:
            row = store._conn.execute(
                "SELECT title, price, want_count FROM item WHERE item_id=?",
                (iid,)).fetchone()
            out.append({"item_id": iid,
                        "title": row["title"] if row else iid,
                        "price": row["price"] if row else None,
                        "want_count": row["want_count"] if row else None,
                        "item_url": f"https://www.goofish.com/item?id={iid}"})
        return out

    return {"available": True, "yesterday": compared_date,
            "new": info(new_ids), "dropped": info(dropped_ids)}


# ---- HTML ----

def _render_html(date_str: str, rdict: dict, store: Store, settings: Settings,
                 extra_notes: list[str] | None = None) -> str:
    stats = store.quality_stats(date_str)
    cmp_result = compare_with_yesterday(
        store, date_str, [s["item_id"] for s in rdict["overall"]],
        Path(settings.paths.data_dir))

    stats_vm = {
        "snapshots": stats.snapshots,
        "items_total": stats.items_total,
        "items_new_today": stats.items_new_today,
        "runs": stats.runs,
        "pages_line": (f"{stats.pages_ok}/{stats.pages_attempted}"
                       if stats.pages_attempted else "—"),
        "completeness": stats.completeness(),
        "pages_degraded": stats.pages_degraded,
        "pages_attempted": stats.pages_attempted,
        "quality_note": ("运行级统计自 run 命令接线后累积"
                         if stats.runs == 0 else ""),
    }
    return _ENV.get_template("report.html.j2").render(
        date=date_str, caliber=rdict["caliber"], stats=stats_vm, cmp=cmp_result,
        overall=rdict["overall"],
        per_keyword=rdict["per_keyword"],
        rising=rdict["rising"], rising_note=rdict["rising_note"],
        extra_notes=extra_notes or [])


# ---- Excel（4 Sheet）----

def _style_header(ws, widths: list[int]) -> None:
    fill = PatternFill("solid", fgColor="1F2329")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _item_row(s: dict, rank=None) -> list:
    return [rank if rank is not None else s.get("rank"),
            s["item_id"], s["title"], s["price"], s["want_count"],
            s.get("delta_want_1d"), s["view_count"],
            round(s["score"], 4) if s.get("score") is not None else None,
            s["seller_nick"], s["location"], s["status"],
            s["source_keyword"], s["item_url"]]


def _write_excel(path: Path, store: Store, date_str: str, rdict: dict) -> None:
    wb = Workbook()

    ws = wb.active
    ws.title = "总榜"
    ws.append(_XLSX_HEADERS)
    for s in rdict["overall"]:
        ws.append(_item_row(s))
    _style_header(ws, _COL_WIDTHS)

    ws = wb.create_sheet("分类榜")
    ws.append(["关键词"] + _XLSX_HEADERS)
    for kw, items in sorted(rdict["per_keyword"].items()):
        for s in items:
            ws.append([kw] + _item_row(s))
    _style_header(ws, [14] + _COL_WIDTHS)

    ws = wb.create_sheet("飙升榜")
    ws.append(_XLSX_HEADERS)
    for s in rdict["rising"]:
        ws.append(_item_row(s))
    if not rdict["rising"]:
        ws.append([rdict["rising_note"]] + [None] * (len(_XLSX_HEADERS) - 1))
    _style_header(ws, _COL_WIDTHS)

    ws = wb.create_sheet("当日全量")
    ws.append(["商品ID", "标题", "价格", "想要数", "状态", "卖家", "所在地",
               "来源关键词", "采集时间", "链接"])
    rows = store._conn.execute(
        """SELECT s.item_id, i.title, s.price, s.want_count, s.status,
                  i.seller_nick, i.location, s.source_keyword, s.crawled_at,
                  i.item_url
           FROM item_snapshot s JOIN item i ON i.item_id = s.item_id
           WHERE s.snapshot_date = ? ORDER BY s.want_count DESC""", (date_str,))
    for row in rows:
        ws.append(list(row))
    _style_header(ws, [16, 46, 9, 9, 8, 14, 10, 14, 20, 40])

    wb.save(path)


# ---- 过期清理（T07 任务 4）----

def cleanup_old_reports(reports_dir: Path, retain_days: int) -> list[str]:
    if not reports_dir.exists():
        return []
    cutoff = (date.today() - timedelta(days=retain_days)).isoformat()
    removed: list[str] = []
    for d in reports_dir.iterdir():
        if d.is_dir() and re.fullmatch(r"\d{4}-\d{2}-\d{2}", d.name) and d.name < cutoff:
            shutil.rmtree(d)
            removed.append(d.name)
    return removed


# ---- 入口 ----

def generate_report(date_str: str, store: Store, settings: Settings,
                    rankings: dict | None = None,
                    extra_notes: list[str] | None = None) -> dict:
    """生成 HTML + Excel 报告。rankings 缺省时从库内现算（T06）。

    extra_notes：数据完整性警示（T11：风控中止/预算用尽等），随口径盒置顶展示。
    """
    cfg = settings.ranking
    full = rankings if rankings is not None else score_day_from_store(
        store, date_str, cfg)
    rdict = rankings_to_dict(full, cfg)

    # 评分结果标准归档（次日对比的数据源）
    parsed_dir = Path(settings.paths.data_dir) / "parsed" / date_str
    parsed_dir.mkdir(parents=True, exist_ok=True)
    (parsed_dir / "_rankings.json").write_text(
        json.dumps(rdict, ensure_ascii=False), encoding="utf-8")

    html = _render_html(date_str, rdict, store, settings, extra_notes=extra_notes)

    report_dir = Path(settings.paths.reports_dir) / date_str
    report_dir.mkdir(parents=True, exist_ok=True)
    html_path = report_dir / "report.html"
    html_path.write_text(html, encoding="utf-8")

    xlsx_path = report_dir / "rankings.xlsx"
    _write_excel(xlsx_path, store, date_str, rdict)

    cleaned = cleanup_old_reports(Path(settings.paths.reports_dir),
                                  settings.report.retain_days)
    return {"html": html_path, "excel": xlsx_path,
            "pool_size": len(full["overall"]), "cleaned": cleaned}
